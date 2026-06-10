#!/usr/bin/env python3
"""
Latitude 43 — city-alert subscriber list, stored in Excel (for now).

One workbook, one sheet, one row per subscriber:

    deals/data/subscribers.xlsx
    | email | name | cities | subscribed_at | updated_at |

`cities` is a comma-joined list of hub slugs (see catalog/_hubs.json).
Re-subscribing with the same email replaces that subscriber's city list,
so the form doubles as a preference-update page.

CLI (also used by whoever sends the emails):

    python3 subscribers.py list --city chicago   # emails to alert for Chicago
    python3 subscribers.py all                   # full table
"""

from __future__ import annotations

import argparse
import os
import re
import sys
from datetime import datetime

from openpyxl import Workbook, load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
DEALS_DIR = os.path.dirname(HERE)
XLSX_PATH = os.path.join(DEALS_DIR, "data", "subscribers.xlsx")
SHEET = "Subscribers"
HEADERS = ["email", "name", "cities", "subscribed_at", "updated_at"]

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def _open() -> tuple:
    if os.path.exists(XLSX_PATH):
        wb = load_workbook(XLSX_PATH)
        ws = wb[SHEET] if SHEET in wb.sheetnames else wb.active
    else:
        os.makedirs(os.path.dirname(XLSX_PATH), exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET
        ws.append(HEADERS)
        # Column widths so the sheet is readable when opened in Excel.
        for col, width in zip("ABCDE", (34, 24, 44, 20, 20)):
            ws.column_dimensions[col].width = width
    return wb, ws


def load_all() -> list[dict]:
    if not os.path.exists(XLSX_PATH):
        return []
    _, ws = _open()
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        rows.append({
            "email": str(row[0]).strip().lower(),
            "name": str(row[1] or "").strip(),
            "cities": [c.strip() for c in str(row[2] or "").split(",") if c.strip()],
            "subscribed_at": str(row[3] or ""),
            "updated_at": str(row[4] or ""),
        })
    return rows


def upsert(email: str, name: str, cities: list[str]) -> dict:
    """Add or update a subscriber. Returns the stored record."""
    email = (email or "").strip().lower()
    if not EMAIL_RE.match(email):
        raise ValueError("invalid email")
    cities = sorted({c.strip() for c in cities if c and c.strip()})
    if not cities:
        raise ValueError("no cities selected")

    wb, ws = _open()
    now = datetime.now().isoformat(timespec="seconds")
    for row in ws.iter_rows(min_row=2):
        if str(row[0].value or "").strip().lower() == email:
            if name:
                row[1].value = name
            row[2].value = ", ".join(cities)
            row[4].value = now
            wb.save(XLSX_PATH)
            return {"email": email, "name": row[1].value, "cities": cities,
                    "updated": True}
    ws.append([email, name, ", ".join(cities), now, now])
    wb.save(XLSX_PATH)
    return {"email": email, "name": name, "cities": cities, "updated": False}


def for_city(city_slug: str) -> list[dict]:
    """Subscribers who asked to hear about this city."""
    return [s for s in load_all() if city_slug in s["cities"]]


def main(argv=None):
    p = argparse.ArgumentParser(description="Latitude 43 subscriber list")
    sub = p.add_subparsers(dest="cmd", required=True)
    pl = sub.add_parser("list", help="Emails subscribed to a city")
    pl.add_argument("--city", required=True, help="hub slug, e.g. chicago")
    sub.add_parser("all", help="Print every subscriber")
    args = p.parse_args(argv)

    if args.cmd == "list":
        subs = for_city(args.city)
        for s in subs:
            print(f"{s['email']}\t{s['name']}")
        print(f"-- {len(subs)} subscriber(s) for {args.city}", file=sys.stderr)
    elif args.cmd == "all":
        for s in load_all():
            print(f"{s['email']}\t{s['name']}\t{', '.join(s['cities'])}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
